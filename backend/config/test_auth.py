from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core import mail
from django.contrib.auth.tokens import default_token_generator
from django.test import Client, TestCase, override_settings
from django.utils.encoding import force_bytes
from django.utils.http import urlsafe_base64_encode
from openpyxl import load_workbook
from rest_framework.authtoken.models import Token

from apps.housing.models import HousingScenario, SavedHousingTest
from finance.models import (
    ExpenseEntry,
    GuestProfile,
    IncomeCoverage,
    IncomeEntry,
    IncomeImportBatch,
    IncomeImportRow,
    UserAppState,
    WorkCostEntry,
)
from finance.services import (
    ensure_default_commitments,
    ensure_default_expense_categories,
    ensure_default_sources,
    ensure_default_work_costs,
)


User = get_user_model()
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_from_response(response):
    return load_workbook(BytesIO(response.content), data_only=True)


def workbook_values(workbook):
    values = []
    for sheet in workbook.worksheets:
        values.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            values.extend("" if value is None else str(value) for value in row)
    return "\n".join(values)


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class AuthApiRegressionTests(TestCase):
    def test_password_reset_response_does_not_reveal_registered_emails(self):
        User.objects.create_user(
            username="known@example.com",
            email="known@example.com",
            password="Passw0rd123",
        )
        client = Client()

        known = client.post(
            "/api/v1/auth/password-reset/",
            data={"email": "known@example.com"},
            content_type="application/json",
        )
        unknown = client.post(
            "/api/v1/auth/password-reset/",
            data={"email": "unknown@example.com"},
            content_type="application/json",
        )

        self.assertEqual(known.status_code, 200)
        self.assertEqual(unknown.status_code, 200)
        self.assertEqual(known.json(), unknown.json())
        self.assertEqual(len(mail.outbox), 1)

    def test_password_reset_resend_sends_another_generic_message(self):
        User.objects.create_user(
            username="resend@example.com",
            email="resend@example.com",
            password="Passw0rd123",
        )
        client = Client()

        first = client.post(
            "/api/v1/auth/password-reset/",
            data={"email": "resend@example.com"},
            content_type="application/json",
        )
        second = client.post(
            "/api/v1/auth/password-reset/",
            data={"email": "resend@example.com"},
            content_type="application/json",
        )

        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(first.json(), second.json())
        self.assertEqual(len(mail.outbox), 2)

    def test_password_reset_confirm_changes_password_and_invalidates_tokens(self):
        user = User.objects.create_user(
            username="reset-confirm@example.com",
            email="reset-confirm@example.com",
            password="OldPassw0rd123",
        )
        token, _ = Token.objects.get_or_create(user=user)
        uid = urlsafe_base64_encode(force_bytes(user.pk))
        reset_token = default_token_generator.make_token(user)
        client = Client(HTTP_AUTHORIZATION=f"Token {token.key}")

        response = client.post(
            "/api/v1/auth/password-reset/confirm/",
            data={"uid": uid, "token": reset_token, "password": "NewPassw0rd123"},
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(Token.objects.filter(user=user).exists())
        old_login = Client().post(
            "/api/v1/auth/login/",
            data={"username": "reset-confirm@example.com", "password": "OldPassw0rd123"},
            content_type="application/json",
        )
        new_login = Client().post(
            "/api/v1/auth/login/",
            data={"username": "reset-confirm@example.com", "password": "NewPassw0rd123"},
            content_type="application/json",
        )
        self.assertEqual(old_login.status_code, 401)
        self.assertEqual(new_login.status_code, 200)

    def test_password_reset_confirm_enforces_password_strength(self):
        user = User.objects.create_user(
            username="reset-strength@example.com",
            email="reset-strength@example.com",
            password="OldPassw0rd123",
        )
        uid = urlsafe_base64_encode(force_bytes(user.pk))
        reset_token = default_token_generator.make_token(user)
        client = Client()

        too_short = client.post(
            "/api/v1/auth/password-reset/confirm/",
            data={"uid": uid, "token": reset_token, "password": "A1short"},
            content_type="application/json",
        )
        no_number = client.post(
            "/api/v1/auth/password-reset/confirm/",
            data={"uid": uid, "token": reset_token, "password": "NoNumberHere"},
            content_type="application/json",
        )

        self.assertEqual(too_short.status_code, 400)
        self.assertEqual(too_short.json()["error"]["code"], "password_too_short")
        self.assertEqual(no_number.status_code, 400)
        self.assertEqual(no_number.json()["error"]["code"], "password_needs_letters_numbers")

    def test_logout_invalidates_the_api_token(self):
        user = User.objects.create_user(
            username="logout@example.com",
            email="logout@example.com",
            password="Passw0rd123",
        )
        token, _ = Token.objects.get_or_create(user=user)
        client = Client(HTTP_AUTHORIZATION=f"Token {token.key}")

        self.assertEqual(client.get("/api/v1/auth/me/").status_code, 200)
        self.assertEqual(client.post("/api/v1/auth/logout/").status_code, 204)
        self.assertEqual(client.get("/api/v1/auth/me/").status_code, 401)

    def test_logout_keeps_authenticated_record_for_next_login(self):
        user = User.objects.create_user(
            username="logout-record@example.com",
            email="logout-record@example.com",
            password="Passw0rd123",
        )
        profile = GuestProfile.objects.create(user=user, session_key="logout-record-profile")
        ensure_default_sources(profile)
        source = profile.income_sources.get(slug="ehail")
        period = profile.financial_periods.create(period_month=date(2026, 4, 1))
        profile.income_entries.create(
            period=period,
            source=source,
            income_date=date(2026, 4, 3),
            gross_amount="1777.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )
        token, _ = Token.objects.get_or_create(user=user)
        client = Client(HTTP_AUTHORIZATION=f"Token {token.key}")

        self.assertEqual(client.post("/api/v1/auth/logout/").status_code, 204)
        login = client.post(
            "/api/v1/auth/login/",
            data={"username": "logout-record@example.com", "password": "Passw0rd123"},
            content_type="application/json",
        )
        self.assertEqual(login.status_code, 200)
        client.defaults["HTTP_AUTHORIZATION"] = f"Token {login.json()['token']}"

        record = client.get("/api/v1/income/record/")

        self.assertEqual(record.status_code, 200)
        self.assertEqual(record.json()["entries"][0]["amount"], "1777.00")

    def test_login_invalid_credentials_are_generic_for_email_or_password(self):
        User.objects.create_user(
            username="login@example.com",
            email="login@example.com",
            password="Passw0rd123",
        )
        client = Client()

        wrong_password = client.post(
            "/api/v1/auth/login/",
            data={"username": "login@example.com", "password": "wrongPass1"},
            content_type="application/json",
        )
        unknown_email = client.post(
            "/api/v1/auth/login/",
            data={"username": "missing@example.com", "password": "wrongPass1"},
            content_type="application/json",
        )

        self.assertEqual(wrong_password.status_code, 401)
        self.assertEqual(unknown_email.status_code, 401)
        self.assertEqual(wrong_password.json(), unknown_email.json())

    def test_login_does_not_silently_claim_guest_record_until_user_accepts(self):
        user = User.objects.create_user(
            username="claim@example.com",
            email="claim@example.com",
            password="Passw0rd123",
        )
        guest_client = Client(HTTP_X_RUMAMPU_CLIENT_ID="claim-guest-client")
        self.assertEqual(guest_client.get("/api/v1/income/record/").status_code, 200)
        guest = GuestProfile.objects.get(user__isnull=True)
        source = guest.income_sources.get(slug="ehail")
        period = guest.financial_periods.create(period_month=date(2026, 5, 1))
        guest.income_entries.create(
            period=period,
            source=source,
            income_date=date(2026, 5, 1),
            gross_amount="2444.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )
        scenario = HousingScenario.objects.create(
            profile=guest,
            property_price="250000.00",
            deposit="25000.00",
            financing_rate="4.000",
            tenure_years=30,
            known_monthly_payment="900.00",
        )

        login = guest_client.post(
            "/api/v1/auth/login/",
            data={"username": "claim@example.com", "password": "Passw0rd123"},
            content_type="application/json",
        )
        self.assertEqual(login.status_code, 200)
        self.assertFalse(GuestProfile.objects.filter(user=user).exists())
        self.assertTrue(GuestProfile.objects.filter(pk=guest.pk, user__isnull=True).exists())
        guest_client.defaults["HTTP_AUTHORIZATION"] = f"Token {login.json()['token']}"

        status_response = guest_client.get("/api/v1/auth/guest-transfer/")
        self.assertEqual(status_response.status_code, 200)
        self.assertTrue(status_response.json()["available"])
        self.assertEqual(status_response.json()["summary"]["income_entries"], 1)

        account_record = guest_client.get("/api/v1/income/record/")
        self.assertEqual(account_record.status_code, 200)
        self.assertEqual(account_record.json()["entries"], [])

        transfer = guest_client.post(
            "/api/v1/auth/guest-transfer/",
            data={"action": "keep"},
            content_type="application/json",
        )
        self.assertEqual(transfer.status_code, 200)
        self.assertTrue(transfer.json()["transferred"])
        self.assertFalse(GuestProfile.objects.filter(pk=guest.pk, user__isnull=True).exists())
        scenario.refresh_from_db()
        self.assertEqual(scenario.user, user)
        self.assertIsNone(scenario.profile_id)

        claimed_record = guest_client.get("/api/v1/income/record/")
        self.assertEqual(claimed_record.status_code, 200)
        self.assertEqual(claimed_record.json()["entries"][0]["amount"], "2444.00")

    def test_declining_guest_transfer_discards_guest_record_and_keeps_account_record(self):
        user = User.objects.create_user(
            username="decline@example.com",
            email="decline@example.com",
            password="Passw0rd123",
        )
        account_profile = GuestProfile.objects.create(user=user, session_key="decline-account")
        ensure_default_sources(account_profile)
        account_source = account_profile.income_sources.get(slug="ehail")
        account_period = account_profile.financial_periods.create(period_month=date(2026, 6, 1))
        account_profile.income_entries.create(
            period=account_period,
            source=account_source,
            income_date=date(2026, 6, 2),
            gross_amount="3000.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )
        guest_client = Client(HTTP_X_RUMAMPU_CLIENT_ID="decline-guest-client")
        self.assertEqual(guest_client.get("/api/v1/income/record/").status_code, 200)
        guest = GuestProfile.objects.get(user__isnull=True)
        guest_source = guest.income_sources.get(slug="ehail")
        guest_period = guest.financial_periods.create(period_month=date(2026, 6, 1))
        guest.income_entries.create(
            period=guest_period,
            source=guest_source,
            income_date=date(2026, 6, 3),
            gross_amount="9999.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )
        login = guest_client.post(
            "/api/v1/auth/login/",
            data={"username": "decline@example.com", "password": "Passw0rd123"},
            content_type="application/json",
        )
        self.assertEqual(login.status_code, 200)
        guest_client.defaults["HTTP_AUTHORIZATION"] = f"Token {login.json()['token']}"

        decline = guest_client.post(
            "/api/v1/auth/guest-transfer/",
            data={"action": "decline"},
            content_type="application/json",
        )

        self.assertEqual(decline.status_code, 200)
        self.assertTrue(decline.json()["discarded"])
        self.assertFalse(GuestProfile.objects.filter(pk=guest.pk).exists())
        record = guest_client.get("/api/v1/income/record/")
        self.assertEqual(record.status_code, 200)
        self.assertEqual(len(record.json()["entries"]), 1)
        self.assertEqual(record.json()["entries"][0]["amount"], "3000.00")

    def test_authenticated_export_returns_readable_xlsx_with_user_data_and_saved_tests(self):
        user = User.objects.create_user(
            username="export@example.com",
            email="export@example.com",
            password="Passw0rd123",
        )
        token, _ = Token.objects.get_or_create(user=user)
        client = Client(HTTP_AUTHORIZATION=f"Token {token.key}")
        profile = GuestProfile.objects.create(user=user, session_key="export-profile")
        ensure_default_sources(profile)
        source = profile.income_sources.get(slug="ehail")
        period = profile.financial_periods.create(period_month=date(2026, 1, 1))
        profile.income_entries.create(
            period=period,
            source=source,
            income_date=date(2026, 1, 1),
            gross_amount="1234.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )
        scenario = HousingScenario.objects.create(
            user=user,
            property_price="300000.00",
            deposit="30000.00",
            financing_rate="4.250",
            tenure_years=30,
            known_monthly_payment="980.00",
        )
        SavedHousingTest.objects.create(
            user=user,
            scenario=scenario,
            name="Apartment near MRT",
            monthly_payment="980.00",
            short_month_count=1,
            tested_months=2,
            largest_gap="120.00",
            income_shock_percent="0.00",
            scenario_snapshot={
                "property_price": "300000.00",
                "deposit": "30000.00",
                "financing_rate": "4.250",
                "tenure_years": 30,
            },
            result_snapshot={
                "months": [{
                    "year": 2026,
                    "month": 1,
                    "gross_income": 1234.0,
                    "usable_income": 1100.0,
                    "existing_costs": 134.0,
                    "available_for_home": 900.0,
                    "tested_home_cost": 980.0,
                    "post_housing_residual": -80.0,
                    "total_shortfall": 80.0,
                    "is_short": True,
                }]
            },
        )

        response = client.get("/api/v1/auth/export/", HTTP_ACCEPT=XLSX_CONTENT_TYPE)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX_CONTENT_TYPE)
        self.assertIn("RuMampu_Record_", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])
        workbook = workbook_from_response(response)
        self.assertEqual(
            workbook.sheetnames,
            ["My Record", "Saved House Tests", "Calculated Results", "About This Export"],
        )
        text = workbook_values(workbook)
        self.assertIn("E-hailing", text)
        self.assertIn("1234", text)
        self.assertIn("Your Data", text)
        self.assertIn("Apartment near MRT", text)
        self.assertIn("300000", text)
        self.assertIn("Calculated", text)
        self.assertIn("2026-01", text)
        self.assertIn("Short", text)

    def test_guest_export_returns_readable_xlsx_without_account_identity(self):
        client = Client(HTTP_X_RUMAMPU_CLIENT_ID="guest-xlsx-export")
        client.get("/api/v1/income/record/")
        profile = GuestProfile.objects.get()
        source = profile.income_sources.get(slug="ehail")
        period = profile.financial_periods.create(period_month=date(2026, 2, 1))
        profile.income_entries.create(
            period=period,
            source=source,
            income_date=date(2026, 2, 1),
            gross_amount="888.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )

        response = client.get("/api/v1/auth/export/", HTTP_ACCEPT=XLSX_CONTENT_TYPE)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX_CONTENT_TYPE)
        workbook = workbook_from_response(response)
        self.assertEqual(
            workbook.sheetnames,
            ["My Record", "Saved House Tests", "Calculated Results", "About This Export"],
        )
        text = workbook_values(workbook)
        self.assertIn("888", text)
        self.assertIn("Your Data", text)
        self.assertNotIn("export@example.com", text)

    def test_export_is_account_isolated_and_omits_internal_sensitive_fields(self):
        user_a = User.objects.create_user(
            username="a-export@example.com",
            email="a-export@example.com",
            password="Passw0rd123",
        )
        user_b = User.objects.create_user(
            username="b-export@example.com",
            email="b-export@example.com",
            password="Passw0rd123",
        )
        profile_a = GuestProfile.objects.create(user=user_a, session_key="profile-a")
        profile_b = GuestProfile.objects.create(user=user_b, session_key="profile-b")
        ensure_default_sources(profile_a)
        ensure_default_sources(profile_b)
        source_a = profile_a.income_sources.get(slug="ehail")
        source_b = profile_b.income_sources.get(slug="ehail")
        period_a = profile_a.financial_periods.create(period_month=date(2026, 3, 1))
        period_b = profile_b.financial_periods.create(period_month=date(2026, 3, 1))
        profile_a.income_entries.create(
            period=period_a,
            source=source_a,
            income_date=date(2026, 3, 1),
            gross_amount="1111.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )
        profile_b.income_entries.create(
            period=period_b,
            source=source_b,
            income_date=date(2026, 3, 1),
            gross_amount="9999.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )
        token, _ = Token.objects.get_or_create(user=user_a)
        client = Client(HTTP_AUTHORIZATION=f"Token {token.key}")

        response = client.get("/api/v1/auth/export/")

        self.assertEqual(response.status_code, 200)
        text = workbook_values(workbook_from_response(response))
        self.assertIn("1111", text)
        self.assertIn("a-export@example.com", text)
        self.assertNotIn("9999", text)
        self.assertNotIn("b-export@example.com", text)
        for internal_label in ("profile_id", "source_id", "category_id", "scenario_id", "slug", "public_id"):
            self.assertNotIn(internal_label, text)
        self.assertNotIn(token.key, text)

    def test_delete_record_removes_authenticated_account_and_token(self):
        user = User.objects.create_user(
            username="delete@example.com",
            email="delete@example.com",
            password="Passw0rd123",
        )
        profile = GuestProfile.objects.create(user=user, session_key="delete-profile")
        ensure_default_sources(profile)
        ensure_default_work_costs(profile)
        ensure_default_commitments(profile)
        ensure_default_expense_categories(profile)
        source = profile.income_sources.get(slug="ehail")
        period = profile.financial_periods.create(period_month=date(2026, 7, 1))
        income = profile.income_entries.create(
            period=period,
            source=source,
            income_date=date(2026, 7, 5),
            gross_amount="1666.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )
        work_cost_item = profile.work_cost_items.get(slug="petrol")
        work_cost = WorkCostEntry.objects.create(
            profile=profile,
            category=work_cost_item,
            cost_date=date(2026, 7, 6),
            amount="55.00",
        )
        commitment = profile.commitment_items.get(slug="rent")
        commitment.monthly_amount = "750.00"
        commitment.save(update_fields=["monthly_amount", "updated_at"])
        expense_category = profile.expense_categories.get(slug="meals")
        expense = ExpenseEntry.objects.create(
            profile=profile,
            category=expense_category,
            expense_date=date(2026, 7, 7),
            amount="18.50",
        )
        import_batch = IncomeImportBatch.objects.create(
            profile=profile,
            file_name="delete-import.csv",
            status=IncomeImportBatch.Status.PREVIEW,
        )
        import_row = IncomeImportRow.objects.create(
            batch=import_batch,
            row_number=1,
            raw_amount="1666",
            raw_date="2026-07-05",
            raw_source="E-hailing",
            amount="1666.00",
            income_date=date(2026, 7, 5),
            source_name="E-hailing",
        )
        coverage = IncomeCoverage.objects.create(
            profile=profile,
            answer=IncomeCoverage.Answer.NO,
        )
        app_state = UserAppState.objects.create(user=user, cash_on_hand="123.00")
        scenario = HousingScenario.objects.create(
            user=user,
            property_price="250000.00",
            deposit="25000.00",
            financing_rate="4.000",
            tenure_years=30,
            known_monthly_payment="900.00",
        )
        SavedHousingTest.objects.create(
            user=user,
            scenario=scenario,
            name="Delete me",
            monthly_payment="900.00",
            short_month_count=0,
            tested_months=1,
            largest_gap="0.00",
        )
        other = User.objects.create_user(
            username="keep-delete@example.com",
            email="keep-delete@example.com",
            password="Passw0rd123",
        )
        other_scenario = HousingScenario.objects.create(
            user=other,
            property_price="260000.00",
            deposit="26000.00",
            financing_rate="4.000",
            tenure_years=30,
            known_monthly_payment="920.00",
        )
        other_saved = SavedHousingTest.objects.create(
            user=other,
            scenario=other_scenario,
            name="Keep me",
            monthly_payment="920.00",
            short_month_count=0,
            tested_months=1,
            largest_gap="0.00",
        )
        token, _ = Token.objects.get_or_create(user=user)
        client = Client(HTTP_AUTHORIZATION=f"Token {token.key}")

        response = client.delete("/api/v1/auth/record/")

        self.assertEqual(response.status_code, 204)
        self.assertFalse(User.objects.filter(email="delete@example.com").exists())
        self.assertFalse(GuestProfile.objects.filter(pk=profile.pk).exists())
        self.assertFalse(UserAppState.objects.filter(pk=app_state.pk).exists())
        self.assertFalse(IncomeEntry.objects.filter(pk=income.pk).exists())
        self.assertFalse(WorkCostEntry.objects.filter(pk=work_cost.pk).exists())
        self.assertFalse(ExpenseEntry.objects.filter(pk=expense.pk).exists())
        self.assertFalse(IncomeImportBatch.objects.filter(pk=import_batch.pk).exists())
        self.assertFalse(IncomeImportRow.objects.filter(pk=import_row.pk).exists())
        self.assertFalse(IncomeCoverage.objects.filter(pk=coverage.pk).exists())
        self.assertFalse(HousingScenario.objects.filter(pk=scenario.pk).exists())
        self.assertFalse(SavedHousingTest.objects.filter(name="Delete me").exists())
        self.assertTrue(User.objects.filter(pk=other.pk).exists())
        self.assertTrue(SavedHousingTest.objects.filter(pk=other_saved.pk).exists())
        self.assertEqual(client.get("/api/v1/auth/me/").status_code, 401)

    def test_deleted_account_credentials_fail_and_email_can_register_again(self):
        user = User.objects.create_user(
            username="recreate@example.com",
            email="recreate@example.com",
            password="Passw0rd123",
        )
        token, _ = Token.objects.get_or_create(user=user)
        client = Client(HTTP_AUTHORIZATION=f"Token {token.key}")
        self.assertEqual(client.delete("/api/v1/auth/record/").status_code, 204)

        old_login = Client().post(
            "/api/v1/auth/login/",
            data={"username": "recreate@example.com", "password": "Passw0rd123"},
            content_type="application/json",
        )
        self.assertEqual(old_login.status_code, 401)

        new_registration = Client().post(
            "/api/v1/auth/register/",
            data={"email": "recreate@example.com", "password": "NewPassw0rd123"},
            content_type="application/json",
        )
        self.assertEqual(new_registration.status_code, 201)
        self.assertTrue(new_registration.json()["token"])

    def test_delete_record_removes_guest_profile_and_housing_scenarios(self):
        client = Client(HTTP_X_RUMAMPU_CLIENT_ID="delete-guest-client")
        self.assertEqual(client.get("/api/v1/income/record/").status_code, 200)
        profile = GuestProfile.objects.get()
        source = profile.income_sources.get(slug="ehail")
        period = profile.financial_periods.create(period_month=date(2026, 8, 1))
        income = profile.income_entries.create(
            period=period,
            source=source,
            income_date=date(2026, 8, 1),
            gross_amount="988.00",
            entry_method=IncomeEntry.EntryMethod.MANUAL,
        )
        work_cost_item = profile.work_cost_items.get(slug="petrol")
        work_cost = WorkCostEntry.objects.create(
            profile=profile,
            category=work_cost_item,
            cost_date=date(2026, 8, 2),
            amount="22.00",
        )
        expense_category = profile.expense_categories.get(slug="meals")
        expense = ExpenseEntry.objects.create(
            profile=profile,
            category=expense_category,
            expense_date=date(2026, 8, 3),
            amount="12.00",
        )
        import_batch = IncomeImportBatch.objects.create(
            profile=profile,
            file_name="guest-delete.csv",
            status=IncomeImportBatch.Status.PREVIEW,
        )
        coverage = IncomeCoverage.objects.create(
            profile=profile,
            answer=IncomeCoverage.Answer.NO,
        )
        scenario = HousingScenario.objects.create(
            profile=profile,
            property_price="180000.00",
            deposit="18000.00",
            financing_rate="4.000",
            tenure_years=30,
            known_monthly_payment="700.00",
        )

        response = client.delete("/api/v1/auth/record/")

        self.assertEqual(response.status_code, 204)
        self.assertFalse(GuestProfile.objects.filter(pk=profile.pk).exists())
        self.assertFalse(IncomeEntry.objects.filter(pk=income.pk).exists())
        self.assertFalse(WorkCostEntry.objects.filter(pk=work_cost.pk).exists())
        self.assertFalse(ExpenseEntry.objects.filter(pk=expense.pk).exists())
        self.assertFalse(IncomeImportBatch.objects.filter(pk=import_batch.pk).exists())
        self.assertFalse(IncomeCoverage.objects.filter(pk=coverage.pk).exists())
        self.assertFalse(HousingScenario.objects.filter(pk=scenario.pk).exists())
