from io import BytesIO
import calendar
import json
import math
import re
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.contrib.auth import authenticate, get_user_model
from django.contrib.auth.tokens import default_token_generator
from django.core.exceptions import ValidationError
from django.core.mail import EmailMultiAlternatives
from django.core.validators import validate_email
from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from django.utils.encoding import force_bytes, force_str
from django.utils.http import urlsafe_base64_decode, urlsafe_base64_encode
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.authtoken.models import Token
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.renderers import BaseRenderer, JSONRenderer
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.housing.models import SavedHousingTest
from finance.models import UserAppState
from finance.services import (
    discard_guest_record_for_request,
    guest_transfer_status,
    profile_for_request,
    transfer_guest_record_to_user,
)


User = get_user_model()
PASSWORD_RESET_PUBLIC_MESSAGE = "If a RuMampu account uses this email, a password reset link has been sent."


def _user_payload(user):
    return {
        "id": user.pk,
        "username": user.get_username(),
        "email": user.email,
    }


def _app_state(user):
    return UserAppState.objects.get_or_create(user=user)[0]


def _auth_payload(user, token=None):
    state = _app_state(user)
    payload = {
        "user": _user_payload(user),
        "cash_on_hand": float(state.cash_on_hand),
        "upfront_costs": state.upfront_costs,
        "docs_checked": state.docs_checked,
        "bought_home": state.bought_home,
        "expense_limits": state.expense_limits,
        "compare_payments": state.compare_payments,
        "saving_plan": state.saving_plan,
        "buffer_state": state.buffer_state,
        "village_state": state.village_state,
        "plan_horizon": state.plan_horizon,
        "pot_moved_months": state.pot_moved_months,
        "kept_tests": state.kept_tests,
        "onboarding_completed": state.onboarding_completed,
        "preferred_language": state.preferred_language,
        "last_record_exported_at": state.last_record_exported_at.isoformat() if state.last_record_exported_at else None,
    }
    if token is not None:
        payload["token"] = token.key
    return payload


_APP_STATE_FIELDS = {
    "cash_on_hand", "upfront_costs", "docs_checked", "bought_home",
    "expense_limits", "compare_payments", "saving_plan", "buffer_state",
    "village_state", "plan_horizon", "pot_moved_months", "kept_tests",
    "onboarding_completed", "preferred_language",
}


def _valid_number(value, *, integer=False, minimum=None):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        return False
    if integer and not isinstance(value, int):
        return False
    return minimum is None or value >= minimum


def _valid_plan(value):
    if value == {}:
        return True
    if not isinstance(value, dict):
        return False
    required = {"key", "target", "n", "amounts", "done", "seed"}
    if set(value) - required - {"skipped", "paused"} or not required <= set(value):
        return False
    key = value["key"]
    match = re.fullmatch(r"(\d{4})-(\d{2})", key) if isinstance(key, str) else None
    if not match:
        return False
    year, month = int(match.group(1)), int(match.group(2))
    if year < 1 or month < 1 or month > 12:
        return False
    n = value["n"]
    if not _valid_number(n, integer=True, minimum=1) or n != calendar.monthrange(year, month)[1]:
        return False
    if not _valid_number(value["target"], minimum=0) or not _valid_number(value["seed"], integer=True, minimum=0):
        return False
    if not isinstance(value["amounts"], list) or len(value["amounts"]) != n:
        return False
    if not all(_valid_number(item, minimum=0) for item in value["amounts"]):
        return False
    if not isinstance(value["done"], list) or len(value["done"]) != n:
        return False
    if not all(isinstance(item, bool) for item in value["done"]):
        return False
    if "skipped" in value and (not isinstance(value["skipped"], list)
                               or len(value["skipped"]) != n
                               or not all(isinstance(item, bool) for item in value["skipped"])):
        return False
    return "paused" not in value or isinstance(value["paused"], bool)


def _valid_state_object(value, allowed, required_arrays=()):
    if not isinstance(value, dict) or set(value) - set(allowed):
        return False
    if value == {}:
        return True
    for key in required_arrays:
        if not isinstance(value.get(key), list):
            return False
    return True


def _valid_buffer_state(value):
    if not _valid_state_object(value, {"saved", "overflow", "target", "houseCost", "prevTarget", "msg"}):
        return False
    if value == {}:
        return True
    for key in ("saved", "overflow"):
        if not _valid_number(value.get(key), minimum=0):
            return False
    for key in ("target", "houseCost", "prevTarget"):
        if value.get(key) is not None and not _valid_number(value.get(key), minimum=0):
            return False
    return value.get("msg") in {None, "used", "moved"}


def _valid_village_state(value):
    allowed = {"cells", "pop", "score", "best", "moves", "gain", "built", "collection", "queued", "savedRm", "msg"}
    if not _valid_state_object(value, allowed, ("cells", "pop")):
        return False
    if value == {}:
        return True
    if len(value["cells"]) != 16 or not all(_valid_number(item, integer=True, minimum=0) and item <= 4 for item in value["cells"]):
        return False
    if not all(_valid_number(item, integer=True, minimum=0) and item < 16 for item in value["pop"]):
        return False
    for key in ("score", "best", "moves", "gain", "built", "collection", "queued", "savedRm"):
        if not _valid_number(value.get(key), minimum=0):
            return False
    return "msg" not in value or isinstance(value["msg"], str)


def _validate_app_state_field(field, value):
    if field == "cash_on_hand":
        if isinstance(value, bool):
            return None
        try:
            amount = Decimal(str(value))
        except (InvalidOperation, ValueError, TypeError):
            return None
        if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
            return None
        return amount
    if field in {"upfront_costs", "docs_checked", "compare_payments", "pot_moved_months", "kept_tests"}:
        if not isinstance(value, list):
            return None
        if field == "docs_checked" and not all(isinstance(item, str) for item in value):
            return None
        if field == "pot_moved_months" and not all(isinstance(item, str) for item in value):
            return None
        if field == "kept_tests":
            for item in value:
                if not isinstance(item, dict) or not all(_valid_number(item.get(key)) for key in ("pay", "s", "n", "g")):
                    return None
        return value
    if field == "expense_limits":
        return value if isinstance(value, dict) else None
    if field == "bought_home":
        return value if isinstance(value, bool) else None
    if field == "plan_horizon":
        return value if value is None or _valid_number(value, integer=True, minimum=1) else None
    if field == "saving_plan":
        return value if _valid_plan(value) else None
    if field == "buffer_state":
        return value if _valid_buffer_state(value) else None
    if field == "village_state":
        return value if _valid_village_state(value) else None
    if field in {"onboarding_completed", "preferred_language"}:
        return value
    return None


def _date(value):
    return value.isoformat() if value else None


def _money(value):
    return float(value)


XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class XlsxRenderer(BaseRenderer):
    media_type = XLSX_CONTENT_TYPE
    format = "xlsx"
    charset = None
    render_style = "binary"

    def render(self, data, accepted_media_type=None, renderer_context=None):
        return data


def _display_money(value):
    if value in (None, ""):
        return None
    return float(value)


def _scenario_value(row, field):
    if row.scenario_id and row.scenario is not None:
        return getattr(row.scenario, field)
    return (row.scenario_snapshot or {}).get(field)


def _style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="D3E7E5")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="243C3E")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in ws.columns:
        max_len = 12
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value) + 2, 42))
        ws.column_dimensions[column_letter].width = max_len
    ws.freeze_panes = "A2"


def _append_rows(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_sheet(ws)


def _build_record_workbook(request, profile):
    now = timezone.localtime()
    workbook = Workbook()

    my_record = workbook.active
    my_record.title = "My Record"
    record_rows = []
    for entry in profile.income_entries.select_related("source").all():
        record_rows.append([
            "Income",
            _date(entry.income_date),
            entry.source.name if entry.source else "Unspecified",
            "",
            _display_money(entry.gross_amount),
            "Your Data",
        ])
    for entry in profile.expense_entries.select_related("category").all():
        record_rows.append([
            "Expense",
            _date(entry.expense_date),
            entry.category.name,
            "",
            _display_money(entry.amount),
            "Your Data",
        ])
    for entry in profile.work_cost_entries.select_related("category").all():
        record_rows.append([
            "Work cost",
            _date(entry.cost_date),
            entry.category.name,
            "",
            _display_money(entry.amount),
            "Your Data",
        ])
    for item in profile.commitment_items.all():
        record_rows.append([
            "Commitment",
            "",
            item.name,
            item.get_commitment_type_display(),
            _display_money(item.monthly_amount),
            "Your Data",
        ])
    _append_rows(
        my_record,
        ["Record type", "Date", "Name / Source / Category", "Category / Type", "Amount (RM)", "Data Type"],
        record_rows,
    )

    saved_tests = workbook.create_sheet("Saved House Tests")
    saved_rows = []
    calculated_rows = []
    if request.user.is_authenticated:
        rows = SavedHousingTest.objects.filter(
            user=request.user,
        ).select_related("scenario").prefetch_related("scenario__additional_costs")
    else:
        rows = SavedHousingTest.objects.none()
    for row in rows:
        name = row.name or "Saved test"
        saved_rows.append([
            name,
            timezone.localtime(row.created_at).strftime("%Y-%m-%d"),
            _display_money(_scenario_value(row, "property_price")),
            _display_money(_scenario_value(row, "deposit")),
            _display_money(_scenario_value(row, "financing_rate")),
            _scenario_value(row, "tenure_years"),
            _display_money(row.monthly_payment),
            row.short_month_count,
            row.tested_months,
            _display_money(row.largest_gap),
            "Your Data and Calculated",
        ])
        result = row.result_snapshot or {}
        for month in result.get("months", []):
            month_label = f"{month.get('year')}-{int(month.get('month', 0)):02d}" if month.get("year") and month.get("month") else ""
            calculated_rows.append([
                name,
                month_label,
                _display_money(month.get("gross_income")),
                _display_money(month.get("usable_income")),
                _display_money(month.get("existing_costs")),
                _display_money(month.get("available_for_home")),
                _display_money(month.get("tested_home_cost")),
                _display_money(month.get("post_housing_residual")),
                _display_money(month.get("total_shortfall", month.get("shortfall"))),
                "Short" if month.get("is_short") else "OK",
                "Calculated",
            ])
    _append_rows(
        saved_tests,
        [
            "Test name", "Test date", "Your Data - Property price (RM)",
            "Your Data - Deposit (RM)", "Your Data - Financing rate (%)",
            "Your Data - Tenure (years)", "Calculated - Tested monthly home cost (RM)",
            "Calculated - Short months", "Calculated - Tested months",
            "Calculated - Largest gap (RM)", "Data Type",
        ],
        saved_rows,
    )

    calculated = workbook.create_sheet("Calculated Results")
    _append_rows(
        calculated,
        [
            "Test name", "Month", "Recorded income (RM)", "Usable income (RM)",
            "Existing costs (RM)", "Available for home (RM)",
            "Tested home cost (RM)", "Amount left after housing (RM)",
            "Shortfall (RM)", "Status", "Data Type",
        ],
        calculated_rows,
    )

    about = workbook.create_sheet("About This Export")
    _append_rows(
        about,
        ["Item", "Details"],
        [
            ["Export generated", now.strftime("%Y-%m-%d %H:%M:%S %Z")],
            ["Signed-in account email", request.user.email if request.user.is_authenticated else ""],
            ["Your Data", "Information you entered or confirmed yourself."],
            ["Calculated", "Figures RuMampu worked out from your record."],
            ["Note", "This file is a personal copy of your RuMampu record and does not depend on RuMampu."],
            ["Privacy", "This export does not include passwords, auth tokens, reset tokens, OAuth secrets, or guest client IDs."],
        ],
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


class RegisterView(APIView):
    """Create a real Django account and authenticate it immediately.

    The UI promises an email-based account with a password of at least eight
    characters containing letters and numbers, so those rules are enforced on
    the server as well as in the React Native form.
    """

    authentication_classes = []
    permission_classes = [AllowAny]

    @transaction.atomic
    def post(self, request):
        email = str(request.data.get("email", "")).strip().lower()
        password = str(request.data.get("password", ""))

        if not email or not password:
            return Response(
                {"error": {"code": "register_required_fields", "message": "Enter your email and password."}},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            validate_email(email)
        except ValidationError:
            return Response(
                {"error": {"code": "invalid_email", "message": "Enter a valid email address."}},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if len(password) < 8:
            return Response(
                {"error": {"code": "password_too_short", "message": "Password must be at least 8 characters."}},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if not any(ch.isalpha() for ch in password) or not any(ch.isdigit() for ch in password):
            return Response(
                {
                    "error": {
                        "code": "password_needs_letters_numbers",
                        "message": "Password must include at least one letter and one number.",
                    }
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # The default Django User model does not require email to be unique, so
        # RuMampu enforces one account per email explicitly. username is also the
        # normalised email so email login remains deterministic.
        if User.objects.filter(email__iexact=email).exists() or User.objects.filter(username__iexact=email).exists():
            return Response(
                {"error": {"code": "account_exists", "message": "An account with this email already exists."}},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # create_user hashes the password; the raw password is never stored.
        user = User.objects.create_user(username=email, email=email, password=password)

        # Guest data is claimed only after explicit consent from the Profile
        # sign-up flow. A normal sign-up starts with a clean account profile.
        merge_guest_data = request.data.get("merge_guest_data", False) is True
        if merge_guest_data:
            original_user = request.user
            request.user = user
            try:
                transfer_guest_record_to_user(request, user)
            finally:
                request.user = original_user
        _app_state(user)

        # Registration also signs the user in. Subsequent API requests use this
        # token and therefore resolve the account-owned profile, not the guest
        # session/profile boundary.
        token, _ = Token.objects.get_or_create(user=user)
        return Response(
            _auth_payload(user, token),
            status=status.HTTP_201_CREATED,
        )


class LoginView(APIView):
    authentication_classes = []
    permission_classes = [AllowAny]

    def post(self, request):
        identifier = str(request.data.get("username", "")).strip()
        password = str(request.data.get("password", ""))

        if not identifier or not password:
            return Response(
                {"error": {"code": "login_required_fields", "message": "Enter your username/email and password."}},
                status=status.HTTP_400_BAD_REQUEST,
            )

        username = identifier
        if "@" in identifier:
            matched_user = User.objects.filter(email__iexact=identifier).only("username").first()
            if matched_user is not None:
                username = matched_user.get_username()

        user = authenticate(request=request, username=username, password=password)
        if user is None or not user.is_active:
            return Response(
                {"error": {"code": "invalid_credentials", "message": "We could not sign you in with those details."}},
                status=status.HTTP_401_UNAUTHORIZED,
            )

        token, _ = Token.objects.get_or_create(user=user)
        return Response(_auth_payload(user, token))


class GuestTransferView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response(guest_transfer_status(request))

    @transaction.atomic
    def post(self, request):
        action = str(request.data.get("action", "")).strip().lower()
        if action == "keep":
            return Response(transfer_guest_record_to_user(request, request.user))
        if action == "decline":
            return Response(discard_guest_record_for_request(request))
        return Response(
            {"error": {"code": "invalid_guest_transfer_action", "message": "Choose whether to keep or discard the guest record."}},
            status=status.HTTP_400_BAD_REQUEST,
        )


class MeView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response(_auth_payload(request.user))

    def patch(self, request):
        state = _app_state(request.user)
        update_fields = []

        try:
            encoded_request = json.dumps(request.data, separators=(",", ":"), allow_nan=False)
        except (TypeError, ValueError):
            return Response(
                {"error": {"code": "invalid_app_state_json", "message": "Account state must be valid JSON."}},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if len(encoded_request.encode("utf-8")) > 32768:
            return Response(
                {"error": {"code": "app_state_too_large", "message": "Account state is too large."}},
                status=status.HTTP_400_BAD_REQUEST,
            )

        for field in _APP_STATE_FIELDS - {"preferred_language", "onboarding_completed"}:
            if field not in request.data:
                continue
            value = _validate_app_state_field(field, request.data[field])
            if value is None and (field != "plan_horizon" or request.data[field] is not None):
                return Response(
                    {"error": {"code": f"invalid_{field}", "message": f"Invalid {field}."}},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            if field == "cash_on_hand":
                changed = state.cash_on_hand != value
            else:
                changed = getattr(state, field) != value
            if changed:
                setattr(state, field, value)
                update_fields.append(field)

        if "preferred_language" in request.data:
            language = str(request.data.get("preferred_language", "")).strip().lower()
            if language not in {"en", "ms", "zh"}:
                return Response(
                    {"error": {"code": "invalid_language", "message": "Choose English, Bahasa Melayu, or Chinese."}},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            if state.preferred_language != language:
                state.preferred_language = language
                update_fields.append("preferred_language")

        if "onboarding_completed" in request.data:
            if request.data.get("onboarding_completed") is not True:
                return Response(
                    {"error": {"code": "invalid_onboarding_state", "message": "onboarding_completed must be true."}},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            if not state.onboarding_completed:
                state.onboarding_completed = True
                update_fields.append("onboarding_completed")

        if not update_fields:
            if not any(key in request.data for key in _APP_STATE_FIELDS):
                return Response(
                    {"error": {"code": "empty_app_state_update", "message": "No supported account settings were provided."}},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            return Response(_auth_payload(request.user))

        update_fields.append("updated_at")
        state.save(update_fields=update_fields)
        return Response(_auth_payload(request.user))


class LogoutView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        Token.objects.filter(user=request.user).delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class RecordExportView(APIView):
    permission_classes = [IsAuthenticated]
    renderer_classes = [XlsxRenderer, JSONRenderer]

    def get(self, request):
        profile = profile_for_request(request)
        filename = timezone.localtime().strftime("RuMampu_Record_%d_%b_%Y.xlsx")
        response = HttpResponse(_build_record_workbook(request, profile), content_type=XLSX_CONTENT_TYPE)
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        state = _app_state(request.user)
        state.last_record_exported_at = timezone.now()
        state.save(update_fields=["last_record_exported_at", "updated_at"])
        return response


class RecordDeleteView(APIView):
    permission_classes = [AllowAny]

    @transaction.atomic
    def delete(self, request):
        profile = profile_for_request(request)
        if request.user.is_authenticated:
            Token.objects.filter(user=request.user).delete()
            request.user.delete()
            return Response(status=status.HTTP_204_NO_CONTENT)
        profile.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class PasswordResetRequestView(APIView):
    """Email a one-time Django password reset token.

    The response is deliberately identical for known and unknown email
    addresses so this endpoint cannot be used to discover registered users.
    """

    authentication_classes = []
    permission_classes = [AllowAny]

    def post(self, request):
        email = str(request.data.get("email", "")).strip().lower()
        if not email:
            return Response(
                {"error": {"code": "reset_email_required", "message": "Enter your email address."}},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            validate_email(email)
        except ValidationError:
            return Response(
                {"error": {"code": "invalid_email", "message": "Enter a valid email address."}},
                status=status.HTTP_400_BAD_REQUEST,
            )

        user = User.objects.filter(email__iexact=email, is_active=True).first()
        if user is None or not user.has_usable_password():
            return Response(
                {"message": PASSWORD_RESET_PUBLIC_MESSAGE},
                status=status.HTTP_200_OK,
            )

        if user.has_usable_password():
            uid = urlsafe_base64_encode(force_bytes(user.pk))
            token = default_token_generator.make_token(user)
            base = str(getattr(settings, "PASSWORD_RESET_URL_BASE", "rumampu://reset-password")).rstrip("?")
            separator = "&" if "?" in base else "?"
            reset_url = f"{base}{separator}uid={uid}&token={token}"

            subject = "Reset your RuMampu password"
            text = (
                "You requested a password reset for your RuMampu account.\n\n"
                f"Open this link to set a new password:\n{reset_url}\n\n"
                "If you did not request this, you can ignore this email."
            )
            html = (
                "<p>You requested a password reset for your RuMampu account.</p>"
                f'<p><a href="{reset_url}">Reset my RuMampu password</a></p>'
                "<p>If you did not request this, you can ignore this email.</p>"
            )
            try:
                message = EmailMultiAlternatives(
                    subject=subject,
                    body=text,
                    from_email=settings.DEFAULT_FROM_EMAIL,
                    to=[user.email],
                )
                message.attach_alternative(html, "text/html")
                message.send(fail_silently=False)
            except Exception:
                return Response(
                    {"message": PASSWORD_RESET_PUBLIC_MESSAGE},
                    status=status.HTTP_200_OK,
                )

        return Response(
            {"message": PASSWORD_RESET_PUBLIC_MESSAGE},
            status=status.HTTP_200_OK,
        )


class PasswordResetConfirmView(APIView):
    """Validate a reset token and replace the user's password."""

    authentication_classes = []
    permission_classes = [AllowAny]

    @transaction.atomic
    def post(self, request):
        uid = str(request.data.get("uid", "")).strip()
        token = str(request.data.get("token", "")).strip()
        password = str(request.data.get("password", ""))

        if not uid or not token or not password:
            return Response(
                {"error": {"code": "reset_required_fields", "message": "The reset link or new password is missing."}},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if len(password) < 8:
            return Response(
                {"error": {"code": "password_too_short", "message": "Password must be at least 8 characters."}},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if not any(ch.isalpha() for ch in password) or not any(ch.isdigit() for ch in password):
            return Response(
                {
                    "error": {
                        "code": "password_needs_letters_numbers",
                        "message": "Password must include at least one letter and one number.",
                    }
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            user_id = force_str(urlsafe_base64_decode(uid))
            user = User.objects.get(pk=user_id, is_active=True)
        except (TypeError, ValueError, OverflowError, User.DoesNotExist):
            user = None

        if user is None or not default_token_generator.check_token(user, token):
            return Response(
                {
                    "error": {
                        "code": "invalid_reset_link",
                        "message": "This password reset link is invalid or has expired. Request a new one.",
                    }
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        user.set_password(password)
        user.save(update_fields=["password"])

        # Invalidate any existing API login tokens. Password reset should log
        # the account out everywhere and require the new password next time.
        Token.objects.filter(user=user).delete()
        return Response({"message": "Your password has been reset. You can now log in with the new password."})
